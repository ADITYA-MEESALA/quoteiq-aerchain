"""Fictional source generation ONLY. This module never produces extracted bids."""
import argparse
import io
import json
import textwrap
from pathlib import Path
import pymupdf as fitz
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from .models import LineItem, Question, RFx, Supplier

BRIEF = "Source 30 corrugated packaging SKUs for an e-commerce fulfilment centre in Bengaluru. Include mailers, shipping cartons, heavy-duty cartons, die-cut trays and partition kits. Require delivery in 21 days, ISO 9001, ECT 32 or better, at least 60% recycled content and sample test reports. Quantities 1,000–8,000 individual pieces. Compare pre-tax delivered costs in INR."


def generate_rfx():
    """Explicit offline seed; AI drafting uses llm.generate_rfx instead."""
    styles = [("E-commerce mailer", "3 ply E flute", "one-colour flexo"),
              ("Shipping carton", "3 ply B flute", "plain kraft"),
              ("Heavy-duty carton", "5 ply BC flute", "two-colour flexo"),
              ("Die-cut display tray", "3 ply B flute", "one-colour flexo"),
              ("Partition kit", "3 ply E flute", "plain kraft")]
    items = []
    for i in range(30):
        label, board, printing = styles[i // 6]
        length, width, height = 180 + (i % 6) * 45, 120 + (i % 6) * 30, 70 + (i % 6) * 25
        items.append(LineItem(id=f"PKG-{i+1:03}", description=f"{label} · {length} × {width} × {height} mm",
                              spec=f"Internal {length}x{width}x{height} mm; {board}; ECT >=32 lb/in; {printing}; recycled fibre >=60%; dimensional tolerance +/-3 mm",
                              quantity=[3000, 5000, 2000, 8000, 1500, 4000][i % 6]))
    return RFx(title="Corrugated packaging · Q4 sourcing", brief=BRIEF, base_currency="INR",
               delivery_location="Bengaluru fulfilment centre, Karnataka", max_lead_days=21,
               commercial_terms="Quote every SKU in individual pieces. State currency, price basis, freight, taxes, discounts and all conditions. Payment requested: net 30 days. Quote validity: 30 days. Delivery in 21 days. Evaluation uses pre-tax delivered cost; no automatic award.",
               items=items, origin="Fictional demo seed (not AI generated)", questionnaire=[
                   Question(id="Q1", question="Do you hold a current ISO 9001 certificate?"),
                   Question(id="Q2", question="Can you certify ECT >=32 lb/in for every quoted SKU?"),
                   Question(id="Q3", question="Is recycled fibre content at least 60%?"),
                   Question(id="Q4", question="Will you provide sample test reports before production?", mandatory=False)])


def demo_suppliers():
    return [Supplier(id=code, name=name, contact=contact, email=email, location=location, notes=notes)
            for code, name, contact, email, location, notes in [
                ("S1", "KraftNest Packaging", "Ananya Rao", "ananya@kraftnest.example", "Bengaluru", "Fictional · regional converter · nonstandard Excel"),
                ("S2", "Deccan Corrupack", "Rohan Desai", "rohan@deccancorrupack.example", "Pune", "Fictional · high-volume mill · PDF with conditional discount"),
                ("S3", "BlueHarbor Cartons", "Maya Chen", "maya@blueharbor.example", "Chennai", "Fictional · export desk · USD Word quotation"),
                ("S4", "GreenFold Industries", "Ishaan Shah", "ishaan@greenfold.example", "Ahmedabad", "Fictional · recycled board · photographed rate card"),
                ("S5", "MetroBox Works", "Farah Khan", "farah@metrobox.example", "Hyderabad", "Fictional · local converter · incomplete informal email"),
            ]]


def price(index, supplier):
    return round((12.4 + index * 1.15 + (index % 4) * 0.85) * [1.02, .98, 1.06, .94, .91][supplier], 2)


def quality_lines(rfx, supplier):
    answers = []
    for idx, q in enumerate(rfx.questionnaire):
        answer = "Yes, we confirm this requirement."
        if supplier == 3 and idx == 1:
            answer = "Pending lab verification; cannot confirm yet."
        if supplier == 4 and idx == 0:
            answer = "No; certification is still in progress."
        answers.append(f"{q.id}: {q.question} Answer: {answer}")
    return answers


def pdf_bytes(lines):
    doc = fitz.open()
    page, y = None, 0
    for line in lines:
        for wrapped in textwrap.wrap(line, 105) or [""]:
            if page is None or y > 790:
                page, y = doc.new_page(width=595, height=842), 45
            page.insert_text((35, y), wrapped, fontsize=9, fontname="helv")
            y += 14
    result = doc.tobytes()
    doc.close()
    return result


def font(size):
    for candidate in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size=size)


def rate_card(rfx):
    # Render real source typography, then simulate a slightly tilted, blurred phone photograph.
    # Values exist only in pixels, with no hidden OCR text or answer sidecar.
    img = Image.new("RGB", (1800, 2600), "#f3efdf")
    draw = ImageDraw.Draw(img)
    draw.rectangle((35, 30, 1760, 2545), outline="#827c69", width=3)
    draw.text((85, 70), "GREENFOLD INDUSTRIES | RATE CARD", font=font(43), fill="#17231d")
    draw.text((85, 132), "Fictional demo source | INR per box | 1 box = 1 piece", font=font(29), fill="#353e37")
    draw.text((85, 195), "SKU                 QUANTITY (pcs)               RATE / BOX", font=font(31), fill="#17231d")
    for i, item in enumerate(rfx.items):
        y = 260 + i * 49
        draw.line((80, y + 42, 1710, y + 42), fill="#c2bdad", width=1)
        draw.text((90, y), item.id, font=font(30), fill="#263829")
        draw.text((540, y), str(item.quantity), font=font(30), fill="#263829")
        draw.text((1130, y), f"INR {price(i, 3):.2f}", font=font(30), fill="#263829")
    foot = ["All technical specifications as per the RFx are accepted.",
            "Freight: INR 3,200 for the full order. No discounts.",
            "Tax excluded from all rates and freight. Delivery: 25 days.",
            "Payment: net 30 days. Validity: 30 days from invitation."] + quality_lines(rfx, 3)
    y = 1790
    for line in foot:
        for wrapped in textwrap.wrap(line, 100):
            draw.text((85, y), wrapped, font=font(25), fill="#374335")
            y += 36
    # A smudge intersects one price and invites evidence-based human review.
    draw.line((1138, 1060, 1310, 1066), fill="#9f9a87", width=4)
    img = img.rotate(-0.7, resample=Image.Resampling.BICUBIC, expand=False, fillcolor="#b8b2a5")
    # Keep the source visibly photographed while retaining enough clarity for a
    # repeatable live multimodal demo. The single smudged rate remains the intended exception.
    img = img.filter(ImageFilter.GaussianBlur(0.35))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=82)
    return out.getvalue()


def save_demo_files(out_dir="demo_files", rfx=None):
    rfx = rfx or generate_rfx()
    root = Path(out_dir)
    root.mkdir(parents=True, exist_ok=True)
    (root / "rfx.json").write_text(rfx.model_dump_json(indent=2), encoding="utf-8")
    paths = {}
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Offer Sep"
    sheet.merge_cells("B2:F2")
    sheet["B2"] = "KRAFTNEST PACKAGING - fictional supplier quotation"
    sheet["B3"] = "Currency INR | rates per 100 pieces | all technical RFx specifications accepted"
    for column, value in enumerate(["Our ref", "Packing article", "Rate / 100 pcs", "Your requirement (pcs)", "Board / finish"], 2):
        cell = sheet.cell(6, column, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="226D58")
    for i, item in enumerate(rfx.items):
        for col, val in enumerate([item.id, item.description, round(price(i, 0) * 100, 2), item.quantity, item.spec], 2):
            sheet.cell(i + 7, col, val)
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 50
    sheet.column_dimensions["D"].width = 22
    sheet.column_dimensions["E"].width = 25
    terms = wb.create_sheet("Small print")
    for line in ["Freight INR 4200 total. No discount offered.", "Tax excluded from all prices and freight.",
                 "Delivery 18 days. Payment net 30. Validity 30 days from invitation."] + quality_lines(rfx, 0):
        terms.append([line])
    path = root / "01_kraftnest_nonstandard.xlsx"
    wb.save(path)
    paths["S1"] = path
    lines = ["DECCAN CORRUPACK | Commercial offer | Fictional demonstration",
             "All rates INR per box. One box means one requested piece. Specifications as per RFx.",
             "SKU | Description | Quantity in pieces | INR per box"]
    for i, item in enumerate(rfx.items):
        lines.append(f"{item.id} | {item.description.replace('×', 'x').replace('·', '-')} | {item.quantity} | {price(i, 1):.2f}")
    lines += ["Freight: INR 6000 for full order. Tax excluded from rates and freight.",
              "Delivery: 20 days. Standard payment: net 30. Validity: 30 days from invitation.",
              "* 3% discount on merchandise ONLY if ALL 30 items are awarded, merchandise subtotal exceeds INR 400000,",
              "and 20% advance is received within 7 days. Freight is not discounted. Otherwise no discount."] + quality_lines(rfx, 1)
    path = root / "02_deccan_conditional.pdf"
    path.write_bytes(pdf_bytes(lines))
    paths["S2"] = path
    doc = Document()
    doc.add_heading("BlueHarbor Cartons", 0)
    doc.add_paragraph("Fictional commercial quotation | USD pricing")
    doc.add_paragraph("Thank you for the opportunity. Our domestic Chennai fulfilment operation quotes in US dollars (USD). All quantities are individual pieces. Every listed article conforms to the corresponding RFx technical specifications. No customs or import charge applies to this fictional domestic delivery.")
    for i, item in enumerate(rfx.items):
        doc.add_paragraph(f"For {item.id}, {item.description}, we can supply {item.quantity:,} pieces at USD {price(i, 2)/85:.4f} per piece. The agreed construction is {item.spec}.")
    doc.add_paragraph("The complete order carries a single freight charge of USD 90. All prices and freight exclude tax. We offer no discount. Delivery is 19 calendar days. Payment is due net 30 days. This offer is valid for 30 days from invitation.")
    doc.add_heading("Quality clarifications", 1)
    for line in quality_lines(rfx, 2):
        doc.add_paragraph(line)
    path = root / "03_blueharbor_usd.docx"
    doc.save(path)
    paths["S3"] = path
    path = root / "04_greenfold_phone_photo.jpg"
    path.write_bytes(rate_card(rfx))
    paths["S4"] = path
    email = ["From: Farah Khan <farah@metrobox.example>", "Subject: box rates, first pass (fictional)", "",
             "Hi team, below is what we can do in INR per box, one box = one piece. These 27 lines meet your technical specs."]
    for i, item in enumerate(rfx.items[:27]):
        email.append(f"{item.id} - {item.quantity} boxes @ INR {price(i, 4):.2f} each")
    email += ["Still checking the last three SKUs. Freight will be confirmed after route planning.",
              "Taxes extra. Payment net 15. We can discuss a discount later; lead time also to be confirmed."] + quality_lines(rfx, 4) + ["Thanks, Farah"]
    path = root / "05_metrobox_email.eml"
    path.write_text("\n".join(email), encoding="utf-8")
    paths["S5"] = path
    (root / "manifest.json").write_text(json.dumps({k: p.name for k, p in paths.items()}, indent=2), encoding="utf-8")
    return paths


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate fictional source files, never extracted answers")
    parser.add_argument("--out", default="demo_files")
    args = parser.parse_args()
    for sid, path in save_demo_files(args.out).items():
        print(f"{sid}: {path}")
