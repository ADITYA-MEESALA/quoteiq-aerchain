"""Format adapters preserve locators and pass images to multimodal extraction."""
import base64
import io
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile
import pymupdf as fitz
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 160_000
SUPPORTED = {".xlsx", ".pdf", ".docx", ".jpg", ".jpeg", ".png", ".txt", ".eml"}


@dataclass
class ParsedDocument:
    text: str
    images: list[tuple[str, bytes]]
    warnings: list[str]

    def content(self):
        content = [{"type": "input_text", "text": self.text or "Read the attached quote image."}]
        for locator, data in self.images:
            content.extend([
                {"type": "input_text", "text": locator},
                {"type": "input_image", "image_url": "data:image/png;base64," + base64.b64encode(data).decode(), "detail": "high"},
            ])
        return content

    def record(self):
        return {"text": self.text, "image_locators": [x[0] for x in self.images], "warnings": self.warnings}


def image_png(data):
    with Image.open(io.BytesIO(data)) as original:
        if original.width * original.height > 25_000_000:
            raise ValueError("Image exceeds 25 megapixels")
        img = ImageOps.exif_transpose(original).convert("RGB")
        img.thumbnail((2400, 3600))
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()


def parse_document(filename, data):
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError("Supported formats: XLSX, PDF, DOCX, JPG, PNG, TXT and EML")
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Source must be nonempty and at most 20 MB")
    if ext in {".xlsx", ".docx"}:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 60 * 1024 * 1024:
                raise ValueError("Expanded office document exceeds 60 MB")
    text, images, warnings = [], [], []
    if ext == ".xlsx":
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        values = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        try:
            for ws in wb:
                if ws.max_row > 1500 or ws.max_column > 60:
                    raise ValueError("Worksheet exceeds 1,500 rows or 60 columns; split the document")
                for row_number, row in enumerate(ws.iter_rows(), 1):
                    cells = []
                    for cell in row:
                        if cell.value is not None:
                            value = str(cell.value)
                            if cell.data_type == "f":
                                cached = values[ws.title][cell.coordinate].value
                                value += f" [cached value: {cached!r}; formula is never executed]"
                                if cached is None:
                                    warnings.append(f"{ws.title}!{cell.coordinate}: formula has no cached result")
                            cells.append(f"{cell.coordinate}={value}")
                    if cells:
                        text.append(f"[Sheet {ws.title}, row {row_number}] " + " | ".join(cells))
        finally:
            wb.close()
            values.close()
    elif ext == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ValueError("Encrypted PDF: upload an unlocked copy")
            if len(doc) > 12:
                raise ValueError("PDF exceeds 12 pages; split it before extraction")
            for n, page in enumerate(doc, 1):
                text.append(f"[Page {n}]\n{page.get_text(sort=True)}")
                pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                images.append((f"Page {n} (rendered source)", pix.tobytes("png")))
    elif ext == ".docx":
        doc = Document(io.BytesIO(data))
        for n, p in enumerate(doc.paragraphs, 1):
            if p.text.strip():
                text.append(f"[Paragraph {n}] {p.text}")
        for t, table in enumerate(doc.tables, 1):
            for r, row in enumerate(table.rows, 1):
                text.append(f"[Table {t}, row {r}] " + " | ".join(c.text for c in row.cells))
        for n, section in enumerate(doc.sections, 1):
            for label, part in (("header", section.header), ("footer", section.footer)):
                for p in part.paragraphs:
                    if p.text:
                        text.append(f"[Section {n} {label}] {p.text}")
        for rel in doc.part.rels.values():
            if "image" in rel.reltype and not rel.is_external:
                try:
                    images.append((f"Embedded image {len(images) + 1}", image_png(rel.target_part.blob)))
                except (OSError, ValueError) as exc:
                    raise ValueError("Unsupported embedded Word image; export to PDF") from exc
        if len(images) > 12:
            raise ValueError("Word document exceeds 12 embedded images")
    elif ext in {".jpg", ".jpeg", ".png"}:
        images.append(("Image 1: uploaded rate card", image_png(data)))
    else:
        decoded = data.decode("utf-8-sig")
        text.extend(f"[Line {i}] {line}" for i, line in enumerate(decoded.splitlines(), 1))
    joined = "\n".join(text)
    if len(joined) > MAX_TEXT:
        raise ValueError("Document exceeds 160,000 text characters; split it before extraction")
    if not joined.strip() and not images:
        raise ValueError("Document contains no readable text or images")
    return ParsedDocument(joined, images, list(dict.fromkeys(warnings)))


def ingest(repo, supplier_id, filename, data, extractor=None):
    """The only ingestion route, shared by demo files, uploads and pasted emails."""
    from .llm import extract_quote, safe_error
    doc_id = repo.add_document(supplier_id, filename, data)
    try:
        parsed = parse_document(filename, data)
        repo.save_parsed(doc_id, parsed.record())
        quote, metadata = (extractor or extract_quote)(repo.rfx(), parsed.content())
        quote.warnings = list(dict.fromkeys(quote.warnings + parsed.warnings))
        repo.finish_document(doc_id, quote.model_dump(), metadata)
    except Exception as exc:
        repo.fail_document(doc_id, safe_error(exc))
    return doc_id
