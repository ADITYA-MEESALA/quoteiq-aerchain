import io
import json
import sqlite3
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from conftest import store_quote
from quoteiq.analyst import AnalystQuery, ask, run_query
from quoteiq.calculations import build_snapshot
from quoteiq.exports import comparison_workbook, csv_bytes, evidence_zip
from quoteiq.extractors import ingest


def test_audit_append_only_and_normalization_idempotent(repo):
    store_quote(repo)
    build_snapshot(repo)
    before = len(repo.audit())
    build_snapshot(repo)
    assert len(repo.audit()) == before
    with pytest.raises(sqlite3.IntegrityError), repo.connection() as con:
        con.execute("DELETE FROM audit")


def test_failed_revision_does_not_replace_active(repo):
    first = store_quote(repo)
    def fail(*args):
        raise ValueError("No result")
    failed = ingest(repo, "S1", "new.txt", b"new source", extractor=fail)
    assert repo.document(failed)["status"] == "failed"
    assert repo.document(first)["active"] == 1


def test_extraction_immutable_and_revision_selection(repo):
    old = store_quote(repo)
    new = store_quote(repo)
    assert repo.document(old)["active"] == 0
    assert repo.document(new)["active"] == 1
    repo.activate_document(old, "Buyer")
    assert repo.document(old)["active"] == 1
    with pytest.raises(ValueError):
        repo.finish_document(old, {}, {})


def test_unknown_or_injected_tool_args_are_rejected(repo):
    snapshot = build_snapshot(repo)
    with pytest.raises(ValueError):
        AnalystQuery(operation="execute_sql", supplier_ids=[], item_ids=[], eligible_only=False)
    with pytest.raises(ValueError):
        AnalystQuery(operation="supplier_totals", supplier_ids=[], item_ids=[], eligible_only=False, sql="DROP TABLE audit")
    with pytest.raises(ValueError):
        run_query({"operation": "supplier_totals", "supplier_ids": ["S999"], "item_ids": [], "eligible_only": False}, snapshot, repo.rfx())
    with pytest.raises(ValueError):
        run_query({"operation": "supplier_totals", "supplier_ids": [], "item_ids": ["PKG-001"], "eligible_only": False}, snapshot, repo.rfx())


def test_tool_grounding_and_scope(repo):
    store_quote(repo)
    snapshot = build_snapshot(repo)
    result = run_query({"operation": "cheapest_eligible_lines", "supplier_ids": ["S1"], "item_ids": ["PKG-001"], "eligible_only": False}, snapshot, repo.rfx())
    assert len(result["records"]) == 1
    assert result["records"][0]["unit_price"] == 10
    assert result["dataset_version"] == snapshot["dataset_version"]
    assert result["unresolved_assumptions"]  # Other suppliers have not quoted.


def test_analyst_function_call_loop_uses_validated_tool(repo):
    store_quote(repo)
    class Responses:
        def __init__(self):
            self.calls = []
        def create(self, **kwargs):
            self.calls.append(kwargs)
            if len(self.calls) == 1:
                call = SimpleNamespace(type="function_call", name="query_quotes", call_id="test-call", arguments=json.dumps({"operation": "supplier_totals", "supplier_ids": [], "item_ids": [], "eligible_only": False}))
                return SimpleNamespace(status="completed", output=[call], output_text="")
            return SimpleNamespace(status="completed", output=[], output_text="The calculation evidence is in [E1]; missing supplier quotes prevent an award decision.")
    responses = Responses()
    result = ask("What can we compare?", build_snapshot(repo), repo.rfx(), api=SimpleNamespace(responses=responses))
    assert result["evidence"][0]["operation"] == "supplier_totals"
    assert responses.calls[0]["tool_choice"] == "required"
    outputs = [x for x in responses.calls[1]["input"] if isinstance(x, dict) and x.get("type") == "function_call_output"]
    assert outputs and "dataset_version" in outputs[0]["output"]


def test_analyst_cannot_answer_without_evidence(repo):
    class Responses:
        def create(self, **kwargs):
            return SimpleNamespace(status="completed", output=[], output_text="A fabricated claim")
    with pytest.raises(RuntimeError, match="no calculation evidence"):
        ask("Cheapest?", build_snapshot(repo), repo.rfx(), api=SimpleNamespace(responses=Responses()))


def test_exports_include_sources_and_escape_spreadsheet_formulas(repo):
    store_quote(repo)
    snapshot = build_snapshot(repo)
    snapshot["suppliers"][0]["supplier"] = "=HYPERLINK(\"evil\")"
    workbook = load_workbook(io.BytesIO(comparison_workbook(snapshot, repo.audit())))
    assert {"Supplier totals", "Normalized lines", "Quality", "Audit", "Assumptions"}.issubset(workbook.sheetnames)
    values = list(workbook["Supplier totals"].values)
    column = values[0].index("supplier")
    assert values[1][column].startswith("'=")
    assert "'=CMD()" in csv_bytes([{"source": "=CMD()"}]).decode("utf-8-sig")
    with ZipFile(io.BytesIO(evidence_zip(repo, snapshot))) as archive:
        assert "snapshot.json" in archive.namelist()
        assert any(x.endswith("test-quote.txt") for x in archive.namelist())
        assert any(x.endswith("reviews.json") for x in archive.namelist())


def test_workspace_reset_clears_demo_records_and_owned_files(repo):
    doc_id = store_quote(repo)
    source = repo.document(doc_id)["path"]
    demo = repo.root / "demo" / "old"
    demo.mkdir(parents=True)
    (demo / "quote.txt").write_text("fictional")

    repo.reset_workspace()

    assert repo.rfx() is None
    assert repo.suppliers() == []
    assert repo.documents() == []
    assert repo.audit() == []
    assert not Path(source).exists()
    assert not (repo.root / "demo").exists()
    repo.log("reset_check", "workspace", reason="Trigger recreated")
    assert len(repo.audit()) == 1
