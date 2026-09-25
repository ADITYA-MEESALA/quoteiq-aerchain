import io
from pathlib import Path
import pymupdf as fitz
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from quoteiq.demo_data import generate_rfx, save_demo_files
from quoteiq.extractors import ingest, parse_document


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    return save_demo_files(tmp_path_factory.mktemp("demo"))


def test_seed_and_five_distinct_real_formats(demo):
    rfx = generate_rfx()
    assert len(rfx.items) == len({i.id for i in rfx.items}) == 30
    assert rfx.origin.startswith("Fictional")
    assert {p.suffix for p in demo.values()} == {".xlsx", ".pdf", ".docx", ".jpg", ".eml"}
    for path in demo.values():
        parsed = parse_document(path.name, path.read_bytes())
        assert parsed.text or parsed.images


def test_nonstandard_excel_locators_and_quantities(demo):
    result = parse_document(demo["S1"].name, demo["S1"].read_bytes())
    assert "B7=PKG-001" in result.text
    assert "100 pcs" in result.text
    assert "Small print" in result.text
    assert "PKG-030" in result.text


def test_pdf_footnote_and_page_images_survive(demo):
    result = parse_document(demo["S2"].name, demo["S2"].read_bytes())
    assert "20% advance" in result.text
    assert "ALL 30" in result.text
    assert "400000" in result.text
    assert result.images and "Page 1" in result.images[0][0]
    assert any(x["type"] == "input_image" for x in result.content())


def test_word_prose_and_usd_survive(demo):
    result = parse_document(demo["S3"].name, demo["S3"].read_bytes())
    assert "USD 90" in result.text
    assert "PKG-030" in result.text
    assert "[Paragraph" in result.text


def test_photo_has_pixels_without_hidden_answer(demo):
    result = parse_document(demo["S4"].name, demo["S4"].read_bytes())
    assert result.text == ""
    assert len(result.images) == 1
    with Image.open(io.BytesIO(result.images[0][1])) as img:
        assert img.width >= 1700 and img.height >= 2500


def test_email_is_27_of_30_and_incomplete(demo):
    result = parse_document(demo["S5"].name, demo["S5"].read_bytes())
    assert "PKG-027" in result.text
    assert "PKG-028" not in result.text
    assert "Freight will be confirmed" in result.text
    assert "[Line 1]" in result.text


def test_formula_is_preserved_not_executed():
    wb = Workbook()
    wb.active["A1"] = "=1+2"
    out = io.BytesIO()
    wb.save(out)
    result = parse_document("test.xlsx", out.getvalue())
    assert "=1+2" in result.text
    assert "formula is never executed" in result.text
    assert result.warnings


def test_docx_tables_headers_and_footers():
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "PKG-009: INR 100"
    doc.sections[0].footer.paragraphs[0].text = "10% subject to full order"
    out = io.BytesIO()
    doc.save(out)
    result = parse_document("quote.docx", out.getvalue())
    assert "[Table 1, row 1]" in result.text
    assert "10% subject to full order" in result.text


def test_png_and_scanned_pdf():
    image = Image.new("RGB", (300, 200), "white")
    png = io.BytesIO()
    image.save(png, "PNG")
    assert parse_document("scan.png", png.getvalue()).images
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_image(page.rect, stream=png.getvalue())
    assert parse_document("scan.pdf", pdf.tobytes()).images
    pdf.close()


@pytest.mark.parametrize("filename,data", [("bad.csv", b"a,b"), ("empty.txt", b""), ("bad.pdf", b"not a pdf")])
def test_invalid_inputs_are_rejected(filename, data):
    with pytest.raises(Exception):
        parse_document(filename, data)


def test_no_key_is_honest_and_original_is_retained(repo, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    data = b"PKG-001: USD 12, freight to be confirmed"
    doc_id = ingest(repo, "S1", "../unsafe.txt", data)
    doc = repo.document(doc_id)
    assert doc["status"] == "failed"
    assert "provider" in doc["error"].lower()  # Check that error mentions provider configuration
    assert doc["filename"] == "unsafe.txt"
    assert Path(doc["path"]).read_bytes() == data
    assert doc["extraction"] is None
    assert doc["parsed"] and doc["sha256"]


def test_all_demo_sources_use_same_injection_boundary(repo, demo):
    from conftest import synthetic_quote
    seen = []
    def extractor(rfx, content):
        seen.append(content)
        return synthetic_quote(rfx), {"test_double": True}
    for sid, path in demo.items():
        doc_id = ingest(repo, sid, path.name, path.read_bytes(), extractor=extractor)
        assert repo.document(doc_id)["status"] == "extracted"
    assert len(seen) == 5
    assert any(p["type"] == "input_image" for p in seen[3])
